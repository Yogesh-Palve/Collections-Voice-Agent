"""
Dual-write persistence: Excel (collections_log.xlsx) + SQLite (collections.db).
"""

from __future__ import annotations

import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from config import Settings, get_settings
from logging_utils import setup_error_logger

logger = setup_error_logger(__name__)

FORWARD_FLOW_COLUMNS = [
    "call_date",
    "call_time",
    "be_name",
    "be_id",
    "be_zone",
    "forward_flow_target_pct",
    "forward_flow_actual_pct",
    "be_was_aware",
    "be_can_reduce",
    "be_strategy",
    "commitment_date",
    "escalated_to_human",
    "call_outcome",
    "notes",
    "session_id",
]

RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

CREATE_TABLE_SQL = """
CREATE TABLE IF NOT EXISTS forward_flow_calls (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    call_date TEXT NOT NULL,
    call_time TEXT NOT NULL,
    be_name TEXT NOT NULL,
    be_id TEXT NOT NULL,
    be_zone TEXT NOT NULL,
    forward_flow_target_pct REAL NOT NULL,
    forward_flow_actual_pct REAL NOT NULL,
    be_was_aware TEXT,
    be_can_reduce TEXT,
    be_strategy TEXT,
    commitment_date TEXT,
    escalated_to_human INTEGER NOT NULL DEFAULT 0,
    call_outcome TEXT NOT NULL,
    notes TEXT,
    session_id TEXT UNIQUE NOT NULL,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
"""


class DataLogger:
    """Excel + SQLite logger for forward flow calls."""

    SHEET_NAME = "ForwardFlow"

    def __init__(self, settings: Optional[Settings] = None) -> None:
        """Ensure data directory and database schema exist."""
        self.settings = settings or get_settings()
        self.data_dir = self.settings.data_path
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.db_path = self.data_dir / "collections.db"
        self.excel_path = self.data_dir / "collections_log.xlsx"
        self._init_db()
        self._init_excel()

    def _init_db(self) -> None:
        """Create SQLite table if missing."""
        with sqlite3.connect(self.db_path) as conn:
            conn.execute(CREATE_TABLE_SQL)
            conn.commit()

    def _init_excel(self) -> None:
        """Create workbook with ForwardFlow sheet and headers if missing."""
        if self.excel_path.exists():
            return
        wb = Workbook()
        ws = wb.active
        ws.title = self.SHEET_NAME
        ws.append(FORWARD_FLOW_COLUMNS)
        wb.save(self.excel_path)

    @staticmethod
    def _format_commitment_date(iso_date: Optional[str]) -> str:
        """Format YYYY-MM-DD as DD-MMM-YYYY for Excel display."""
        if not iso_date:
            return ""
        try:
            d = datetime.strptime(iso_date[:10], "%Y-%m-%d").date()
            return d.strftime("%d-%b-%Y")
        except ValueError:
            return iso_date

    def log_forward_flow_call(self, record: dict) -> None:
        """
        Append one forward-flow call to Excel and SQLite.

        Args:
            record: Dict matching ForwardFlowCallRecord fields.
        """
        row = {k: record.get(k, "") for k in FORWARD_FLOW_COLUMNS}
        commitment_display = self._format_commitment_date(
            record.get("commitment_date") or None
        )
        row["commitment_date"] = commitment_display or record.get("commitment_date", "")

        try:
            self._append_sqlite(record)
        except Exception as exc:
            logger.error("SQLite log failed: %s", exc)

        try:
            self._append_excel(row, record)
        except Exception as exc:
            logger.error("Excel log failed: %s", exc)

    def _append_sqlite(self, record: dict) -> None:
        """Insert row into forward_flow_calls."""
        with sqlite3.connect(self.db_path) as conn:
            conn.execute(
                """
                INSERT OR REPLACE INTO forward_flow_calls (
                    call_date, call_time, be_name, be_id, be_zone,
                    forward_flow_target_pct, forward_flow_actual_pct,
                    be_was_aware, be_can_reduce, be_strategy, commitment_date,
                    escalated_to_human, call_outcome, notes, session_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    record.get("call_date"),
                    record.get("call_time"),
                    record.get("be_name"),
                    record.get("be_id"),
                    record.get("be_zone"),
                    record.get("forward_flow_target_pct"),
                    record.get("forward_flow_actual_pct"),
                    record.get("be_was_aware", ""),
                    record.get("be_can_reduce", ""),
                    record.get("be_strategy", ""),
                    record.get("commitment_date"),
                    1 if record.get("escalated_to_human") else 0,
                    record.get("call_outcome", "In Progress"),
                    record.get("notes", ""),
                    record.get("session_id"),
                ),
            )
            conn.commit()

    def _append_excel(self, row: dict, raw_record: dict) -> None:
        """Append row to ForwardFlow sheet with conditional highlighting."""
        wb = load_workbook(self.excel_path)
        ws = wb[self.SHEET_NAME]
        values = [row.get(col, "") for col in FORWARD_FLOW_COLUMNS]
        ws.append(values)
        row_idx = ws.max_row

        target = float(raw_record.get("forward_flow_target_pct") or 0)
        actual = float(raw_record.get("forward_flow_actual_pct") or 0)
        can_reduce = str(raw_record.get("be_can_reduce", "")).strip()

        fill = None
        if can_reduce in ("No", "Unsure"):
            fill = RED_FILL
        elif actual > target + 3:
            fill = YELLOW_FILL

        if fill:
            for col in range(1, len(FORWARD_FLOW_COLUMNS) + 1):
                ws.cell(row=row_idx, column=col).fill = fill

        wb.save(self.excel_path)

    def get_calls_by_be(self, be_id: str) -> list[dict]:
        """Fetch all forward-flow calls for a Business Executive."""
        with sqlite3.connect(self.db_path) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT * FROM forward_flow_calls WHERE be_id = ? ORDER BY call_date DESC",
                (be_id,),
            ).fetchall()
        return [dict(r) for r in rows]

    def get_pending_followups(self, as_of_date: Optional[date] = None) -> list[dict]:
        """
        Calls where commitment_date <= as_of_date and outcome is Committed.

        Args:
            as_of_date: Defaults to today.
        """
        as_of = as_of_date or date.today()
        as_of_str = as_of.isoformat()
        with sqlite3.connect(self.db_path) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                """
                SELECT * FROM forward_flow_calls
                WHERE call_outcome = 'Committed'
                  AND commitment_date IS NOT NULL
                  AND commitment_date <= ?
                ORDER BY commitment_date ASC
                """,
                (as_of_str,),
            ).fetchall()
        return [dict(r) for r in rows]

    def get_calls_in_range(self, start: date, end: date) -> list[dict]:
        """Fetch calls between start and end dates (inclusive)."""
        with sqlite3.connect(self.db_path) as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                """
                SELECT * FROM forward_flow_calls
                WHERE call_date >= ? AND call_date <= ?
                ORDER BY call_date DESC, call_time DESC
                """,
                (start.isoformat(), end.isoformat()),
            ).fetchall()
        return [dict(r) for r in rows]

    def export_summary_report(self, month: str) -> Path:
        """
        Export monthly summary to a new sheet in the Excel workbook.

        Args:
            month: YYYY-MM format.

        Returns:
            Path to the Excel file.
        """
        year, mon = map(int, month.split("-"))
        start = date(year, mon, 1)
        if mon == 12:
            end = date(year + 1, 1, 1)
        else:
            end = date(year, mon + 1, 1)
        from datetime import timedelta

        end = end - timedelta(days=1)

        calls = self.get_calls_in_range(start, end)
        wb = load_workbook(self.excel_path)
        sheet_name = f"Summary_{month.replace('-', '_')}"
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)

        ws.append(["Metric", "Value"])
        ws.append(["Month", month])
        ws.append(["Total Calls", len(calls)])
        committed = sum(1 for c in calls if c.get("call_outcome") == "Committed")
        escalated = sum(1 for c in calls if c.get("call_outcome") == "Escalated")
        ws.append(["Committed", committed])
        ws.append(["Escalated", escalated])
        ws.append(["No Answer", sum(1 for c in calls if c.get("call_outcome") == "No Answer")])

        ws.append([])
        ws.append(["BE Name", "BE ID", "Target %", "Actual %", "Outcome"])
        for c in calls:
            ws.append(
                [
                    c.get("be_name"),
                    c.get("be_id"),
                    c.get("forward_flow_target_pct"),
                    c.get("forward_flow_actual_pct"),
                    c.get("call_outcome"),
                ]
            )

        for col in range(1, 6):
            ws.column_dimensions[get_column_letter(col)].width = 18

        wb.save(self.excel_path)
        return self.excel_path


# Module-level convenience functions
_default_logger: Optional[DataLogger] = None


def _get_logger() -> DataLogger:
    global _default_logger
    if _default_logger is None:
        _default_logger = DataLogger()
    return _default_logger


def log_forward_flow_call(record: dict) -> None:
    """Write forward-flow call to Excel and SQLite."""
    _get_logger().log_forward_flow_call(record)


def get_calls_by_be(be_id: str) -> list[dict]:
    """Fetch all calls for a BE."""
    return _get_logger().get_calls_by_be(be_id)


def get_pending_followups(as_of_date: Optional[date] = None) -> list[dict]:
    """Fetch commitment follow-ups due on or before as_of_date."""
    return _get_logger().get_pending_followups(as_of_date)


def export_summary_report(month: str) -> Path:
    """Export monthly summary sheet."""
    return _get_logger().export_summary_report(month)
